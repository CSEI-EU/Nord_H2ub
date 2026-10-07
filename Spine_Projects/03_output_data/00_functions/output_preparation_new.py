# this is going to be new output preparation script 
"""
Recreates the 'calculations' sheet from the raw model-output workbook,
using pandas + numpy instead of Excel formulas.

Raw sheets required (all assumed present, unmodified, in the input file):
    - report__model
    - report__node__stochastic_scenar
    - report__unit__stochastic_scenar
    - Inv_cost           (has a lifetime-in-years formula in col G, but we
                           recompute it from col F ourselves, so a cached
                           value is not required)

Reproduces, cell-for-cell, the logic of the original 'calculations' sheet:

    Storage Investment block (methanol / hydrogen / power):
        total_share_of_max = SUM(report__node__stochastic_scenar!<col>26282:26293)
        total_cost         = total_share_of_max * Inv_cost.unit_cost * n_horizons * Inv_cost.lifetime_years
        O&M                = total_cost * Inv_cost.om_rate

    Unit Investment block (13 units):
        total      = SUM(report__unit__stochastic_scenar!<col>2:13)
        total_cost = total * Inv_cost.unit_cost * n_horizons * Inv_cost.lifetime_years   (only for 6 units)
        O&M        = total_cost * Inv_cost.om_rate                                        (only for 6 units)

NOTE on a quirk in the original spreadsheet, preserved here on purpose:
    The methanol storage 'total_cost' formula multiplies by Inv_cost row 1
    (solar_plant's unit_cost) instead of row 15 (ch3oh_st's own unit_cost),
    while its lifetime/O&M still come from row 15. This looks like a
    copy-paste bug in the original file, but since the goal is to
    reproduce the workbook's actual output, it is kept as-is below
    (flagged with STORAGE_UNIT_COST_ROW).
"""

import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _col_sum(ws, col_letter: str, row_start: int, row_end: int) -> float:
    """Sum a single-column range, e.g. col_sum(ws, 'F', 26282, 26293)."""
    col = column_index_from_string(col_letter)
    vals = [ws.cell(row=r, column=col).value for r in range(row_start, row_end + 1)]
    vals = [v for v in vals if v is not None]
    return float(np.sum(vals)) if vals else 0.0


def _parse_lifetime_days(raw) -> float:
    """'12775D' -> 12775.0   (mirrors Excel's VALUE(LEFT(text,LEN(text)-1)))"""
    if raw is None:
        return np.nan
    m = re.match(r"\s*([\d.]+)", str(raw))
    return float(m.group(1)) if m else np.nan


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------

def load_inv_cost(wb) -> pd.DataFrame:
    """Inv_cost sheet -> DataFrame indexed by the original Excel row number (1-16)."""
    ws = wb["Inv_cost"]
    rows = []
    for r in range(1, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if name is None:
            continue
        rows.append(
            {
                "excel_row": r,
                "name": name,
                "unit_cost": ws.cell(row=r, column=5).value,       # col E
                "lifetime_days_raw": ws.cell(row=r, column=6).value,  # col F, e.g. '12775D'
                "om_rate": ws.cell(row=r, column=8).value,          # col H
            }
        )
    df = pd.DataFrame(rows).set_index("excel_row")
    df["lifetime_years"] = df["lifetime_days_raw"].apply(_parse_lifetime_days) / 365
    return df


def count_rolling_horizons(wb) -> int:
    """Rolling_Horizons!B2 = COUNTIF(report__model!C2:C49, 'total_costs')"""
    ws = wb["report__model"]
    metrics = [ws.cell(row=r, column=3).value for r in range(2, 50)]
    return int(sum(1 for m in metrics if m == "total_costs"))


# ---------------------------------------------------------------------------
# Main calculation
# ---------------------------------------------------------------------------

# Storage technologies: label -> (report__node__stochastic_scenar column, Inv_cost row for lifetime/O&M)
STORAGE_MAP = {
    "methanol": {"data_col": "F", "inv_row": 15},  # ch3oh_st
    "hydrogen": {"data_col": "J", "inv_row": 14},  # h2_st
    "power":    {"data_col": "O", "inv_row": 16},  # power_st
}
# Quirk preserved from the original workbook: methanol's total_cost formula
# multiplies by row 1 (solar_plant) instead of row 15 (ch3oh_st).
STORAGE_UNIT_COST_ROW = {"methanol": 1, "hydrogen": 14, "power": 16}

# Units: report__unit__stochastic_scenar column -> unit name
UNIT_COLUMNS = {
    "D": "ch3oh_reactor",
    "E": "co2_import",
    "F": "co2_vaporizer",
    "G": "dh_heat_exchanger",
    "H": "dist_tower",
    "I": "electrolyzer",
    "J": "excess_heat_exchanger",
    "K": "o2_liquefier",
    "L": "pth_dummy_unit",
    "M": "solar_plant",
    "N": "steam_plant",
    "O": "water_import",
    "P": "wind_plant",
}
# Only these 6 units had a total_cost/O&M formula in the original sheet
# (the rest were left blank there, so we leave them blank here too).
UNIT_INV_ROW = {
    "co2_vaporizer": 3,
    "dist_tower": 4,
    "electrolyzer": 2,
    "excess_heat_exchanger": 11,
    "o2_liquefier": 13,
    "steam_plant": 6,
}


def compute_calculations(path: str):
    wb = load_workbook(path, data_only=True)

    inv_cost = load_inv_cost(wb)
    n_horizons = count_rolling_horizons(wb)

    ws_node = wb["report__node__stochastic_scenar"]
    ws_unit = wb["report__unit__stochastic_scenar"]

    # --- Storage Investment block ---
    storage_rows = []
    for label, cfg in STORAGE_MAP.items():
        total_share = _col_sum(ws_node, cfg["data_col"], 26282, 26293)
        inv_row = inv_cost.loc[cfg["inv_row"]]
        unit_cost_row = inv_cost.loc[STORAGE_UNIT_COST_ROW[label]]

        total_cost = total_share * unit_cost_row["unit_cost"] * n_horizons * inv_row["lifetime_years"]
        om = total_cost * inv_row["om_rate"]

        storage_rows.append(
            {"label": label, "total_share_of_max": total_share, "total_cost": total_cost, "O&M": om}
        )
    storage_df = pd.DataFrame(storage_rows).set_index("label")

    # --- Unit Investment block ---
    unit_rows = []
    for col, name in UNIT_COLUMNS.items():
        total = _col_sum(ws_unit, col, 2, 13)
        row = {"unit": name, "total": total, "total_cost": np.nan, "O&M": np.nan}
        if name in UNIT_INV_ROW:
            inv_row = inv_cost.loc[UNIT_INV_ROW[name]]
            total_cost = total * inv_row["unit_cost"] * n_horizons * inv_row["lifetime_years"]
            row["total_cost"] = total_cost
            row["O&M"] = total_cost * inv_row["om_rate"]
        unit_rows.append(row)
    unit_df = pd.DataFrame(unit_rows).set_index("unit")

    return {
        "n_horizons": n_horizons,
        "inv_cost": inv_cost,
        "storage_investment": storage_df,
        "unit_investment": unit_df,
    }


if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "C:/Users/kar.eco/Downloads/example_output.xlsx"
    result = compute_calculations(path)

    print(f"Rolling horizons (n): {result['n_horizons']}\n")
    print("=== Storage Investment ===")
    print(result["storage_investment"].to_string())
    print("\n=== Unit Investment ===")
    print(result["unit_investment"].to_string())